"""文档加载：把 PDF / Word / TXT / Markdown 解析成纯文本。"""

from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def _load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 同时抽取表格里的文字（教务通知/评定办法常用表格）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_document(path: str | Path) -> str:
    """根据扩展名解析文档，返回纯文本。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"不支持的文件类型: {ext}。支持: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        text = _load_pdf(path)
    elif ext == ".docx":
        text = _load_docx(path)
    else:  # .txt / .md / .markdown
        text = _load_text(path)

    text = text.strip()
    if not text:
        raise ValueError(f"未能从 {path.name} 中提取到任何文本（可能是扫描版 PDF）。")
    return text
